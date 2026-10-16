"""Artifact export implementations for JSON, Markdown, DOCX, XLSX and ReqIF."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Tuple
from xml.etree import ElementTree as ET

from backend.repository import ArtifactRecord


_MIME = {
    "json": "application/json",
    "md": "text/markdown; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "reqif": "application/xml",
    "puml": "text/plain; charset=utf-8",
}


def _requirements(record: ArtifactRecord) -> list[dict]:
    if not record.structured_data:
        return []
    return list(record.structured_data.get("requirements") or [])


def _trace_rows(record: ArtifactRecord) -> list[dict]:
    rows = []
    for requirement in _requirements(record):
        steps = requirement.get("source_steps") or []
        ids = requirement.get("source_step_ids") or []
        locations = requirement.get("source_locations") or []
        for index, step_id in enumerate(ids):
            location = locations[index] if index < len(locations) else {}
            rows.append({
                "requirement_id": requirement.get("id", ""),
                "type": requirement.get("type", ""),
                "text": requirement.get("text", ""),
                "source_feature": requirement.get("source_feature", ""),
                "source_scenario": requirement.get("source_scenario", ""),
                "source_step_id": step_id,
                "source_step": steps[index] if index < len(steps) else "",
                "source_file": location.get("source_file", ""),
                "start_line": location.get("start_line", ""),
                "end_line": location.get("end_line", ""),
                "priority": requirement.get("priority", ""),
                "confidence": requirement.get("confidence", ""),
            })
    return rows


def export_json(record: ArtifactRecord) -> bytes:
    return json.dumps(record.to_dict(), ensure_ascii=False, indent=2).encode("utf-8")


def export_markdown(record: ArtifactRecord) -> bytes:
    return record.content.encode("utf-8")


def export_csv(record: ArtifactRecord) -> bytes:
    output = io.StringIO(newline="")
    rows = _trace_rows(record)
    fields = list(rows[0].keys()) if rows else ["artifact_id", "output_type", "content"]
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    if rows:
        writer.writerows(rows)
    else:
        writer.writerow({"artifact_id": record.id, "output_type": record.output_type, "content": record.content})
    return output.getvalue().encode("utf-8-sig")


def export_docx(record: ArtifactRecord) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading("BDD-to-Requirements Artifact", level=0)
    document.add_paragraph(f"Artifact ID: {record.id}")
    document.add_paragraph(f"Type: {record.output_type}")
    document.add_paragraph(f"Model: {record.provider}/{record.model}")
    document.add_paragraph(f"Prompt version: {record.prompt_version}")
    document.add_paragraph(f"Status: {record.status}")

    requirements = _requirements(record)
    if requirements:
        document.add_heading("Requirements", level=1)
        for requirement in requirements:
            document.add_heading(requirement.get("id", "Requirement"), level=2)
            document.add_paragraph(requirement.get("text", ""))
            table = document.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for label, value in (
                ("Type", requirement.get("type", "")),
                ("Priority", requirement.get("priority", "")),
                ("Confidence", requirement.get("confidence", "")),
                ("Feature", requirement.get("source_feature", "")),
                ("Scenario", requirement.get("source_scenario", "")),
            ):
                row = table.add_row().cells
                row[0].text = str(label)
                row[1].text = str(value)
            document.add_paragraph("Source steps:", style="List Bullet")
            for step_id, step_text in zip(
                requirement.get("source_step_ids") or [],
                requirement.get("source_steps") or [],
            ):
                document.add_paragraph(f"{step_id}: {step_text}", style="List Bullet 2")
            document.add_paragraph("Acceptance criteria:", style="List Bullet")
            for criterion in requirement.get("acceptance_criteria") or []:
                document.add_paragraph(str(criterion), style="List Bullet 2")
    else:
        document.add_heading("Generated Content", level=1)
        for line in record.content.splitlines():
            document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_xlsx(record: ArtifactRecord) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    requirements_sheet = workbook.active
    requirements_sheet.title = "Requirements"
    headers = [
        "ID", "Type", "Requirement", "Priority", "Confidence", "Feature",
        "Scenario", "Acceptance Criteria", "Parameters", "Warnings",
    ]
    requirements_sheet.append(headers)
    for cell in requirements_sheet[1]:
        cell.font = Font(bold=True)
    for requirement in _requirements(record):
        requirements_sheet.append([
            requirement.get("id", ""),
            requirement.get("type", ""),
            requirement.get("text", ""),
            requirement.get("priority", ""),
            requirement.get("confidence", ""),
            requirement.get("source_feature", ""),
            requirement.get("source_scenario", ""),
            "\n".join(requirement.get("acceptance_criteria") or []),
            ", ".join(requirement.get("parameters") or []),
            "\n".join(requirement.get("warnings") or []),
        ])
    requirements_sheet.freeze_panes = "A2"
    requirements_sheet.auto_filter.ref = requirements_sheet.dimensions
    for column, width in {"A": 14, "B": 20, "C": 70, "D": 12, "E": 12, "F": 35, "G": 35, "H": 60, "I": 25, "J": 45}.items():
        requirements_sheet.column_dimensions[column].width = width

    trace_sheet = workbook.create_sheet("Traceability")
    trace_rows = _trace_rows(record)
    trace_headers = list(trace_rows[0].keys()) if trace_rows else ["No traceability data"]
    trace_sheet.append(trace_headers)
    for cell in trace_sheet[1]:
        cell.font = Font(bold=True)
    for row in trace_rows:
        trace_sheet.append([row[key] for key in trace_headers])
    trace_sheet.freeze_panes = "A2"

    metadata_sheet = workbook.create_sheet("Metadata")
    for key, value in (
        ("Artifact ID", record.id),
        ("Output type", record.output_type),
        ("Provider", record.provider),
        ("Model", record.model),
        ("Prompt version", record.prompt_version),
        ("Created at", record.created_at),
        ("Status", record.status),
        ("Review note", record.review_note),
    ):
        metadata_sheet.append([key, str(value)])
    for cell in metadata_sheet["A"]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_reqif(record: ArtifactRecord) -> bytes:
    """Export a self-contained ReqIF 1.2 document.

    The document declares its datatype, attribute definitions, object type,
    specification type, objects and hierarchy, avoiding dangling references.
    """
    ns = "http://www.omg.org/spec/ReqIF/20110401/reqif.xsd"
    ET.register_namespace("", ns)

    def node(parent, tag, **attrs):
        return ET.SubElement(parent, f"{{{ns}}}{tag}", attrs)

    root = ET.Element(f"{{{ns}}}REQ-IF")
    header = node(root, "THE-HEADER")
    reqif_header = node(header, "REQ-IF-HEADER", IDENTIFIER=f"HEADER-{record.id}")
    node(reqif_header, "CREATION-TIME").text = datetime.now(timezone.utc).isoformat()
    node(reqif_header, "REQ-IF-TOOL-ID").text = "BDD-to-Requirements"
    node(reqif_header, "REQ-IF-VERSION").text = "1.2"
    node(reqif_header, "SOURCE-TOOL-ID").text = f"{record.provider}/{record.model}"
    node(reqif_header, "TITLE").text = f"BDD Requirements {record.id}"

    core = node(root, "CORE-CONTENT")
    content = node(core, "REQ-IF-CONTENT")

    datatypes = node(content, "DATATYPES")
    node(
        datatypes,
        "DATATYPE-DEFINITION-STRING",
        IDENTIFIER="DT-STRING",
        **{"LONG-NAME": "Text", "MAX-LENGTH": "32768"},
    )

    spec_types = node(content, "SPEC-TYPES")
    object_type = node(
        spec_types,
        "SPEC-OBJECT-TYPE",
        IDENTIFIER="TYPE-REQUIREMENT",
        **{"LONG-NAME": "Requirement"},
    )
    spec_attributes = node(object_type, "SPEC-ATTRIBUTES")
    attribute_names = {
        "ATTR-TEXT": "Requirement Text",
        "ATTR-TYPE": "Requirement Type",
        "ATTR-PRIORITY": "Priority",
        "ATTR-SOURCE": "Traceability Source",
    }
    for identifier, long_name in attribute_names.items():
        definition = node(
            spec_attributes,
            "ATTRIBUTE-DEFINITION-STRING",
            IDENTIFIER=identifier,
            **{"LONG-NAME": long_name},
        )
        attr_type = node(definition, "TYPE")
        node(attr_type, "DATATYPE-DEFINITION-STRING-REF").text = "DT-STRING"

    node(
        spec_types,
        "SPECIFICATION-TYPE",
        IDENTIFIER="TYPE-SPECIFICATION",
        **{"LONG-NAME": "Requirements Specification"},
    )

    spec_objects = node(content, "SPEC-OBJECTS")
    object_ids: list[str] = []
    for requirement in _requirements(record):
        object_id = requirement.get("id") or uuid_fallback(requirement)
        object_ids.append(object_id)
        spec_object = node(
            spec_objects,
            "SPEC-OBJECT",
            IDENTIFIER=object_id,
            **{"LONG-NAME": object_id},
        )
        values = node(spec_object, "VALUES")
        source_parts = []
        locations = requirement.get("source_locations") or []
        for index, step_id in enumerate(requirement.get("source_step_ids") or []):
            location = locations[index] if index < len(locations) else {}
            source_parts.append(
                f"{step_id} {location.get('source_file', '')}:"
                f"{location.get('start_line', '')}-{location.get('end_line', '')}"
            )
        attributes = {
            "ATTR-TEXT": requirement.get("text", ""),
            "ATTR-TYPE": requirement.get("type", ""),
            "ATTR-PRIORITY": requirement.get("priority", ""),
            "ATTR-SOURCE": "; ".join(source_parts),
        }
        for definition_id, value_text in attributes.items():
            value = node(values, "ATTRIBUTE-VALUE-STRING", **{"THE-VALUE": str(value_text)})
            definition = node(value, "DEFINITION")
            node(definition, "ATTRIBUTE-DEFINITION-STRING-REF").text = definition_id
        type_node = node(spec_object, "TYPE")
        node(type_node, "SPEC-OBJECT-TYPE-REF").text = "TYPE-REQUIREMENT"

    node(content, "SPEC-RELATIONS")
    specifications = node(content, "SPECIFICATIONS")
    specification = node(
        specifications,
        "SPECIFICATION",
        IDENTIFIER=f"SPEC-{record.id}",
        **{"LONG-NAME": f"BDD Requirements {record.id}"},
    )
    spec_type = node(specification, "TYPE")
    node(spec_type, "SPECIFICATION-TYPE-REF").text = "TYPE-SPECIFICATION"
    children = node(specification, "CHILDREN")
    for index, object_id in enumerate(object_ids, 1):
        hierarchy = node(
            children,
            "SPEC-HIERARCHY",
            IDENTIFIER=f"HIER-{record.id}-{index:04d}",
        )
        obj = node(hierarchy, "OBJECT")
        node(obj, "SPEC-OBJECT-REF").text = object_id

    node(content, "SPEC-RELATION-GROUPS")
    node(root, "TOOL-EXTENSIONS")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def uuid_fallback(requirement: dict) -> str:
    digest = hashlib.sha256(str(requirement.get("text", "")).encode("utf-8")).hexdigest()[:16]
    return "REQ-" + digest.upper()


def export_artifact(record: ArtifactRecord, format_name: str) -> Tuple[bytes, str, str]:
    format_name = format_name.lower().strip().lstrip(".")
    if format_name == "markdown":
        format_name = "md"
    if record.output_type == "uml" and format_name in {"md", "txt"}:
        format_name = "puml"
    exporters = {
        "json": export_json,
        "md": export_markdown,
        "puml": export_markdown,
        "csv": export_csv,
        "docx": export_docx,
        "xlsx": export_xlsx,
        "reqif": export_reqif,
    }
    if format_name not in exporters:
        raise ValueError("Unsupported export format")
    data = exporters[format_name](record)
    filename = f"{record.output_type}_{record.id}.{format_name}"
    return data, _MIME.get(format_name, "application/octet-stream"), filename
